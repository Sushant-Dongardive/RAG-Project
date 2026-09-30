import re
from pathlib import Path
import pypdf
import pdfplumber
import pandas as pd


def clean_text(text: str) -> str:
    if not text:
        return ""
    text = re.sub(r'(\w+)-\s*\n\s*(\w+)', r'\1\2', text)
    text = re.sub(r'[^\S\r\n]+', ' ', text)
    text = re.sub(r'\n{3,}', '\n\n', text)
    return text.strip()


def is_toc_page(text: str) -> bool:
    if not text:
        return False
    if len(re.findall(r"\.{4,}", text)) >= 3:
        return True
    first_lines = "\n".join(text.split("\n")[:5]).lower()
    toc_indicators = ["table of contents", "index", "list of figures", "list of tables"]
    return any(ind in first_lines for ind in toc_indicators)


def extract_tables_as_markdown(page) -> str:
    try:
        tables = page.extract_tables()
        if not tables:
            return ""
        md_tables = []
        for table in tables:
            if not table or len(table) < 2:
                continue
            header = [str(cell or "").strip().replace("\n", " ") for cell in table[0]]
            if not any(header):
                continue
            divider = ["---"] * len(header)
            rows = []
            for row in table[1:]:
                clean_row = [str(cell or "").strip().replace("\n", " ") for cell in row]
                rows.append("| " + " | ".join(clean_row) + " |")
            table_md = "\n| " + " | ".join(header) + " |\n| " + " | ".join(divider) + " |\n" + "\n".join(rows) + "\n"
            md_tables.append(table_md)
        return "\n".join(md_tables)
    except Exception:
        return ""


def extract_text_from_pdf(pdf_path: str):
    pages_data = []
    # Validate magic bytes before feeding to PDF engines
    try:
        with open(pdf_path, "rb") as f:
            header = f.read(4)
        if header.startswith(b"PK"):
            # ZIP/Docx file disguised as PDF
            return extract_text_from_docx(pdf_path)
    except Exception:
        pass

    try:
        with pdfplumber.open(pdf_path) as pdf:
            for page_idx, page in enumerate(pdf.pages):
                page_text = page.extract_text() or ""
                if is_toc_page(page_text):
                    continue
                cleaned = clean_text(page_text)
                table_md = extract_tables_as_markdown(page)
                if table_md:
                    cleaned = cleaned + "\n\n### Data Table:\n" + table_md
                if len(cleaned.strip()) > 20:
                    pages_data.append({"page": page_idx + 1, "text": cleaned})
    except Exception as e:
        print(f"[Chunker Warning] pdfplumber error on {pdf_path}: {e}")

    if not pages_data:
        try:
            reader = pypdf.PdfReader(pdf_path)
            for idx, page in enumerate(reader.pages):
                raw = page.extract_text() or ""
                if is_toc_page(raw):
                    continue
                extracted = clean_text(raw)
                if len(extracted) > 20:
                    pages_data.append({"page": idx + 1, "text": extracted})
        except Exception as e:
            print(f"[Chunker Error] pypdf error on {pdf_path}: {e}")

    return pages_data


def extract_text_from_docx(docx_path: str):
    try:
        import docx
        doc = docx.Document(docx_path)
        blocks = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            rows_data = []
            for row in table.rows:
                cells = [c.text.strip().replace("\n", " ") for c in row.cells]
                if any(cells):
                    rows_data.append("| " + " | ".join(cells) + " |")
            if len(rows_data) > 1:
                header = rows_data[0]
                divider = "| " + " | ".join(["---"] * len(table.rows[0].cells)) + " |"
                blocks.append("\n" + header + "\n" + divider + "\n" + "\n".join(rows_data[1:]) + "\n")
        full_doc = "\n\n".join(blocks)
        return [{"page": 1, "text": full_doc}] if len(full_doc) > 20 else []
    except Exception as e:
        print(f"[Chunker Warning] docx extraction failed: {e}")
        return []


def extract_text_from_csv(csv_path: str):
    try:
        df = pd.read_csv(csv_path)
        md = df.to_markdown(index=False)
        return [{"page": 1, "text": md}] if md else []
    except Exception as e:
        print(f"[Chunker Warning] CSV extraction failed on {csv_path}: {e}")
        return []


def extract_text_from_plain(file_path: str):
    try:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            content = clean_text(f.read())
        return [{"page": 1, "text": content}] if len(content) > 10 else []
    except Exception:
        return []


def extract_document(file_path: str):
    ext = Path(file_path).suffix.lower()
    if ext == ".pdf":
        return extract_text_from_pdf(file_path)
    elif ext in [".docx", ".doc"]:
        return extract_text_from_docx(file_path)
    elif ext == ".csv":
        return extract_text_from_csv(file_path)
    elif ext in [".txt", ".md", ".json", ".log"]:
        return extract_text_from_plain(file_path)
    return []


def recursive_chunk_text(text: str, chunk_size: int = 500, overlap: int = 100):
    if not text:
        return []
    blocks = [b.strip() for b in text.split("\n\n") if b.strip()]
    chunks = []
    current_blocks = []
    current_words = 0

    for block in blocks:
        block_words = len(block.split())
        if current_words + block_words > chunk_size and current_blocks:
            chunks.append("\n\n".join(current_blocks))
            current_blocks = [current_blocks[-1]] if len(current_blocks) > 1 else []
            current_words = sum(len(b.split()) for b in current_blocks)
        current_blocks.append(block)
        current_words += block_words

    if current_blocks:
        chunks.append("\n\n".join(current_blocks))

    seen = set()
    unique = []
    for c in chunks:
        key = c[:100].strip()
        if key not in seen and len(c.strip()) > 30:
            seen.add(key)
            unique.append(c)
    return unique